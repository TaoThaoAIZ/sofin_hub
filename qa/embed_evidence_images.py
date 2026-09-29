# -*- coding: utf-8 -*-
"""Nhúng thẳng ảnh evidence (e2e/evidence/*.png) vào 1 sheet riêng trong SofinHub_TestCases.xlsx,
để xem được ở BẤT KỲ đâu (Excel local, Google Drive/Sheets...) — không phụ thuộc link file:// trên máy.
Chạy SAU khi đã chạy qa/apply_playwright_results.py.
"""
import json
import os
import re

import openpyxl
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from PIL import Image as PILImage

XLSX = r"C:\Users\dev\Desktop\sofin_hub\SofinHub_TestCases.xlsx"
RESULTS = r"C:\Users\dev\Desktop\sofin_hub\e2e\parsed_results.json"
EVIDENCE_DIR = r"C:\Users\dev\Desktop\sofin_hub\e2e\evidence"
THUMB_DIR = r"C:\Users\dev\Desktop\sofin_hub\e2e\evidence\_thumbs"
FONT_NAME = "Arial"

THUMB_WIDTH = 360  # px

os.makedirs(THUMB_DIR, exist_ok=True)

with open(RESULTS, encoding="utf-8") as f:
    results = json.load(f)

status_vn = {"passed": "Pass", "failed": "Fail", "skipped": "N/A"}
status_fill = {
    "Pass": PatternFill("solid", fgColor="C6EFCE"),
    "Fail": PatternFill("solid", fgColor="FFC7CE"),
    "N/A": PatternFill("solid", fgColor="D9D9D9"),
}

wb = openpyxl.load_workbook(XLSX)
if "Bằng chứng (ảnh)" in wb.sheetnames:
    del wb["Bằng chứng (ảnh)"]
ws = wb.create_sheet("Bằng chứng (ảnh)")

header_fill = PatternFill("solid", fgColor="1F4E78")
header_font = Font(name=FONT_NAME, size=10, bold=True, color="FFFFFF")
headers = ["Mã Test Case", "Tiêu đề", "Trạng thái test 1", "Ảnh chụp (Playwright)"]
for i, h in enumerate(headers, start=1):
    c = ws.cell(row=1, column=i, value=h)
    c.font = header_font
    c.fill = header_fill
    c.alignment = Alignment(horizontal="center", vertical="center")
ws.row_dimensions[1].height = 22
ws.freeze_panes = "A2"

col_widths = [16, 46, 14, 54]
for i, w in enumerate(col_widths, start=1):
    ws.column_dimensions[get_column_letter(i)].width = w

ROW_PT_PER_PX = 0.75  # xấp xỉ quy đổi pixel ảnh -> point chiều cao dòng Excel

row = 2
embedded = 0
for tc_id, r in sorted(results.items()):
    src = os.path.join(EVIDENCE_DIR, f"{tc_id}.png")
    ws.cell(row=row, column=1, value=tc_id).font = Font(name=FONT_NAME, size=10, bold=True)
    title = re.sub(r"^TC-[A-Z]+-\d+:\s*", "", r["title"])
    ws.cell(row=row, column=2, value=title).font = Font(name=FONT_NAME, size=9)
    ws.cell(row=row, column=2).alignment = Alignment(wrap_text=True, vertical="top")
    status = status_vn.get(r["status"], "N/A")
    sc = ws.cell(row=row, column=3, value=status)
    sc.font = Font(name=FONT_NAME, size=10, bold=True)
    sc.fill = status_fill.get(status)
    sc.alignment = Alignment(horizontal="center", vertical="center")

    if os.path.exists(src):
        thumb_path = os.path.join(THUMB_DIR, f"{tc_id}.jpg")
        with PILImage.open(src) as im:
            im = im.convert("RGB")
            w0, h0 = im.size
            th_h = int(h0 * (THUMB_WIDTH / w0))
            im = im.resize((THUMB_WIDTH, th_h), PILImage.LANCZOS)
            im.save(thumb_path, "JPEG", quality=72, optimize=True)
        xl_img = XLImage(thumb_path)
        xl_img.width = THUMB_WIDTH
        xl_img.height = th_h
        anchor = f"D{row}"
        ws.add_image(xl_img, anchor)
        ws.row_dimensions[row].height = max(th_h * ROW_PT_PER_PX, 18)
        embedded += 1
    else:
        nc = ws.cell(row=row, column=4, value="(không có ảnh — test bị bỏ qua hoàn toàn, xem e2e/report/index.html)")
        nc.font = Font(name=FONT_NAME, size=9, italic=True, color="808080")
        ws.row_dimensions[row].height = 18
    row += 1

ws.auto_filter.ref = f"A1:D{row - 1}"

wb.save(XLSX)
size_mb = os.path.getsize(XLSX) / (1024 * 1024)
print(f"Embedded {embedded} images. New xlsx size: {size_mb:.1f} MB")
