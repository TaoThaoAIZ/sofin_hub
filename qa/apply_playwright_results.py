# -*- coding: utf-8 -*-
"""Ghi kết quả chạy Playwright (e2e/parsed_results.json) vào cột Test 1 của SofinHub_TestCases.xlsx."""
import json
import re
import openpyxl

ANSI_RE = re.compile(r'\x1b\[[0-9;]*m')
ILLEGAL_XLSX_RE = re.compile(r'[\x00-\x08\x0b\x0c\x0e-\x1f]')

def clean(text: str) -> str:
    text = ANSI_RE.sub('', text)
    text = ILLEGAL_XLSX_RE.sub('', text)
    return text.strip()

import os
from openpyxl.styles import Font

XLSX = r"C:\Users\dev\Desktop\sofin_hub\SofinHub_TestCases.xlsx"
RESULTS = r"C:\Users\dev\Desktop\sofin_hub\e2e\parsed_results.json"
EVIDENCE_DIR = r"C:\Users\dev\Desktop\sofin_hub\e2e\evidence"
REPORT_FILE = r"C:\Users\dev\Desktop\sofin_hub\e2e\report\index.html"
REPORT_LINK = "e2e/report/index.html"
TESTER1 = "Claude (Playwright automation)"
HYPERLINK_FONT = Font(name="Arial", size=10, color="0563C1", underline="single")

def to_file_url(path: str) -> str:
    return "file:///" + path.replace("\\", "/")

STATUS_MAP = {"passed": "Pass", "failed": "Fail", "skipped": "N/A"}

# Ghi chú thủ công cho các case đã điều chỉnh cách test so với mô tả gốc trong BRD/test case,
# hoặc case bị skip nhưng lý do không nằm trong annotation runtime.
MANUAL_NOTES = {
    "TC-HOME-009": "ĐÃ ĐIỀU CHỈNH: UI chỉ có dropdown loại giá (Miễn phí/Có phí/Dùng thử), không có ô nhập khoảng giá min-max như test case gốc mô tả — test theo dropdown 'Có phí' thay thế.",
    "TC-HOME-028": "ĐÃ ĐIỀU CHỈNH: ứng dụng không đồng bộ phân trang lên URL (state cục bộ trong useCourseFilters, không dùng useSearchParams) nên không thể test bằng ?page=999 — chuyển sang kiểm tra nút 'Trang sau' tự vô hiệu hóa ở trang cuối.",
    "TC-HOME-029": "ĐÃ ĐIỀU CHỈNH: tương tự TC-HOME-028 — kiểm tra nút 'Trang trước' tự vô hiệu hóa ở trang 1 thay vì ?page=0/-1.",
    "TC-AUTH-032": "BỎ QUA: REFRESH_TOKEN_TTL_DAYS=30 theo cấu hình — không thể chờ thật 30 ngày trong 1 lần chạy. Cần môi trường test riêng với TTL rút ngắn hoặc test-hook cấp sẵn refresh token hết hạn.",
}

with open(RESULTS, encoding="utf-8") as f:
    results = json.load(f)

wb = openpyxl.load_workbook(XLSX)
# Testcase nằm ở nhiều sheet (mỗi module 1 sheet); nhận diện bằng ô A1 = "Mã Test Case" và B1 = "Module" (sheet Bằng chứng có A1 giống nhưng B1 = "Tiêu đề").
TEST_SHEETS = [w for w in wb.worksheets if w.cell(row=1, column=1).value == "Mã Test Case" and w.cell(row=1, column=2).value == "Module"]
updated = 0
for ws in TEST_SHEETS:

    header = [c.value for c in ws[1]]
    col = {name: i + 1 for i, name in enumerate(header)}

    for row in ws.iter_rows(min_row=2):
        tc_id = row[0].value
        if tc_id not in results:
            continue
        r = results[tc_id]
        status = STATUS_MAP.get(r["status"], "N/A")
        ws.cell(row=row[0].row, column=col["Người test 1 (Automation)"], value=TESTER1)
        ws.cell(row=row[0].row, column=col["Trạng thái test 1"], value=status)

        screenshot_path = os.path.join(EVIDENCE_DIR, f"{tc_id}.png")
        ev_cell = ws.cell(row=row[0].row, column=col["Evidences"])
        if os.path.exists(screenshot_path):
            # Xem trực tiếp trong sheet "Bằng chứng (ảnh)" (nhúng sẵn, mở được cả trên Google Drive/Sheets).
            # Link file:// dưới đây CHỈ hoạt động khi mở file Excel thật trên máy này (không dùng được qua
            # trình duyệt/Google Drive vì trình duyệt không có quyền đọc ổ đĩa cục bộ).
            ev_cell.value = f"📷 Xem ảnh ở sheet 'Bằng chứng (ảnh)' — hoặc bấm đây nếu mở Excel local: {tc_id}.png"
            ev_cell.hyperlink = to_file_url(screenshot_path)
        else:
            ev_cell.value = f"Không có ảnh riêng (test bị bỏ qua) — trace đầy đủ: {REPORT_LINK} (chỉ mở được khi tải file .xlsx và mở Excel local)"
            ev_cell.hyperlink = to_file_url(REPORT_FILE)
        ev_cell.font = HYPERLINK_FONT

        # Cột Ghi chú ở đây HOÀN TOÀN do script này sinh ra (ghi đè, không cộng dồn) — nếu người dùng tự
        # gõ ghi chú thủ công vào ô này, hãy chuyển sang cột khác trước khi chạy lại script.
        notes = []
        if tc_id in MANUAL_NOTES:
            notes.append(MANUAL_NOTES[tc_id])
        if r["status"] == "skipped" and r["error"]:
            notes.append(f"Bỏ qua: {clean(r['error'])}")
        for a in r["annotations"]:
            if a:
                notes.append(clean(a))
        if r["status"] == "failed" and r["error"]:
            notes.append(f"Lỗi assertion: {clean(r['error'])}")
        ws.cell(row=row[0].row, column=col["Ghi chú / Kết quả thực tế / Bug ref"], value=" || ".join(notes) if notes else None)
        updated += 1


wb.save(XLSX)
print(f"Updated {updated} rows")
