# -*- coding: utf-8 -*-
"""Tách sheet tạm "Test Cases" (1 sheet duy nhất, ~1.7K dòng) thành MỖI MODULE MỘT SHEET để dễ nhìn/lọc/giao việc.

gen_testcases.py vẫn dựng dữ liệu vào 1 sheet tạm (giữ nguyên cột A..R để các script khác như
apply_playwright_results.py dùng chung), sau đó gọi `split_test_cases(...)` để:
  * tạo sheet riêng cho từng module (tên "<MÃ> - <tên ngắn>"), sao chép nguyên giá trị/định dạng/hyperlink,
    freeze panes, autofilter, dropdown trạng thái;
  * xóa sheet tạm;
  * viết lại công thức ở sheet "Tổng quan" (đếm trên từng sheet module) + thêm cột tiến độ test (Pass/Fail/Chưa test)
    và liên kết bấm nhảy tới từng sheet.
"""
import re
from copy import copy

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

FONT_NAME = "Arial"

# Tên ngắn (<= 31 ký tự tính cả "<MÃ> - "). Mã lạ -> lấy phần đầu của tên module.
SHORT_NAMES = {
    "HOME": "Trang chủ & Khám phá",
    "AUTH": "Tài khoản & Xác thực",
    "COMM": "Cộng đồng",
    "COMMVP": "Cộng đồng MVP",
    "COURSE": "Lớp học",
    "CERT": "Chứng nhận",
    "FEED": "Bảng tin",
    "MEMBER": "Thành viên & Xếp hạng",
    "EVENT": "Lịch & Sự kiện",
    "NOTI": "Thông báo & Nhắn tin",
    "SEARCH": "Tìm kiếm",
    "UPLOAD": "Upload tệp & ảnh",
    "PAY": "Thanh toán & Gói",
    "ADMIN": "Quản trị & Kiểm duyệt",
    "ROLE": "Ma trận phân quyền",
    "SEC": "Bảo mật & Phi chức năng",
    "INTEG": "API & Concurrency",
}
# Thứ tự hiển thị các sheet module (theo luồng người dùng); mã không có trong danh sách xếp cuối theo thứ tự xuất hiện.
ORDER = ["HOME", "AUTH", "COMM", "COMMVP", "COURSE", "CERT", "FEED", "MEMBER", "EVENT", "NOTI", "SEARCH", "UPLOAD",
         "PAY", "ADMIN", "ROLE", "SEC", "INTEG"]

_BAD = set('[]:*?/\\')


def sheet_title(code, name):
    short = SHORT_NAMES.get(code) or name
    t = f"{code} - {short}"
    t = "".join(ch for ch in t if ch not in _BAD)
    return t[:31]


def _q(title):
    return "'" + title.replace("'", "''") + "'"


def split_test_cases(wb, stage, overview, seen_modules, data_start, hdr_row, done_label, plan_label):
    """Trả về danh sách [(code, name, sheet_title)] theo thứ tự đã tạo."""
    header = [c.value for c in stage[1]]
    ncols = len(header)
    header_cells = list(stage[1])

    # gom dòng theo mã module (đọc từ tiền tố mã TC: TC-<MÃ>-<số>)
    by_code = {}
    names = {}
    for row in stage.iter_rows(min_row=2):
        tc = row[0].value
        if not tc:
            continue
        code = tc.split("-")[1]
        by_code.setdefault(code, []).append(row)
        names.setdefault(code, row[1].value)

    codes = [c for c in ORDER if c in by_code] + [c for c in by_code if c not in ORDER]
    stage_idx = wb.sheetnames.index(stage.title)

    widths = {k: v.width for k, v in stage.column_dimensions.items()}
    created = []
    for n, code in enumerate(codes):
        rows = by_code[code]
        title = sheet_title(code, names[code])
        ws = wb.create_sheet(title, index=stage_idx + 1 + n)
        # header
        for ci, hc in enumerate(header_cells, start=1):
            c = ws.cell(row=1, column=ci, value=hc.value)
            c._style = copy(hc._style)
        ws.row_dimensions[1].height = 32
        # dữ liệu
        for ri, row in enumerate(rows, start=2):
            for ci in range(ncols):
                src = row[ci]
                dst = ws.cell(row=ri, column=ci + 1, value=src.value)
                dst._style = copy(src._style)
                if src.hyperlink is not None and getattr(src.hyperlink, "target", None):
                    dst.hyperlink = src.hyperlink.target
            h = stage.row_dimensions[row[0].row].height
            if h:
                ws.row_dimensions[ri].height = h
        for k, w in widths.items():
            ws.column_dimensions[k].width = w
        # Trang này chỉ có 1 module: cột "Module" thừa -> thu hẹp (không ẩn để khỏi lệch chỉ số cột của các script khác).
        ws.column_dimensions["B"].width = 12
        last = len(rows) + 1
        ws.freeze_panes = "E2"
        ws.auto_filter.ref = f"A1:{get_column_letter(ncols)}{last}"
        ws.sheet_properties.tabColor = "1F4E78" if code not in ("SEC", "ROLE", "INTEG", "ADMIN") else "C00000"
        for col_letter, formula in (
            ("N", '"Chưa test,Pass,Fail,Blocked,N/A"'),
            ("Q", '"Chưa test,Pass,Fail,Blocked,N/A"'),
            ("F", '"Cao,Trung bình,Thấp"'),
            ("G", f'"{done_label},{plan_label}"'),
            ("H", '"Có,Một phần,Không"'),
        ):
            dv = DataValidation(type="list", formula1=formula, allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"{col_letter}2:{col_letter}{last}")
        created.append((code, names[code], title))

    del wb[stage.title]

    # ------------------------------------------------------------------ Tổng quan
    sheets = [t for _, _, t in created]
    by_title = {code: t for code, _, t in created}

    def total(col, crit):
        return "+".join(f"COUNTIF({_q(t)}!${col}:${col},{crit})" for t in sheets)

    normal = Font(name=FONT_NAME, size=10)
    link = Font(name=FONT_NAME, size=10, color="0563C1", underline="single")
    hfont = Font(name=FONT_NAME, size=10, bold=True, color="FFFFFF")
    hfill = PatternFill("solid", fgColor="1F4E78")

    # Cột tiến độ thêm sau 7 cột sẵn có
    extra_headers = ["Test 1: Pass", "Test 2: Pass", "Fail (T1+T2)", "Chưa test (T1)", "Chưa test (T2)"]
    for i, h in enumerate(extra_headers):
        c = overview.cell(row=hdr_row, column=8 + i, value=h)
        c.font, c.fill, c.alignment = hfont, hfill, Alignment(horizontal="center")
        overview.column_dimensions[get_column_letter(8 + i)].width = 15

    for i, (code, name) in enumerate(seen_modules):
        rr = data_start + i
        t = by_title.get(code)
        if not t:
            continue
        q = _q(t)
        cell = overview.cell(row=rr, column=2, value=name)
        cell.hyperlink = f"#{q}!A1"
        cell.font = link
        overview.cell(row=rr, column=1, value=code).font = normal
        f = {
            3: f"=COUNTA({q}!$A$2:$A$5000)",
            4: f'=COUNTIF({q}!$G:$G,"{done_label}")',
            5: f'=COUNTIF({q}!$G:$G,"{plan_label}")',
            6: f'=COUNTIF({q}!$H:$H,"Có")',
            7: f"=C{rr}-F{rr}",
            8: f'=COUNTIF({q}!$N:$N,"Pass")',
            9: f'=COUNTIF({q}!$Q:$Q,"Pass")',
            10: f'=COUNTIF({q}!$N:$N,"Fail")+COUNTIF({q}!$Q:$Q,"Fail")',
            11: f'=COUNTIF({q}!$N:$N,"Chưa test")',
            12: f'=COUNTIF({q}!$Q:$Q,"Chưa test")',
        }
        for col, formula in f.items():
            overview.cell(row=rr, column=col, value=formula).font = normal

    total_row = data_start + len(seen_modules)
    for col in range(8, 13):
        L = get_column_letter(col)
        c = overview.cell(row=total_row, column=col, value=f"=SUM({L}{data_start}:{L}{total_row - 1})")
        c.font = Font(name=FONT_NAME, size=10, bold=True)

    # Các khối thống kê (Trạng thái/Loại test/Ưu tiên/Playwright): cộng dồn trên mọi sheet module
    pat = re.compile(r"^=COUNTIF\('Test Cases'!\$([A-Z]):\$[A-Z],(\w+)\)$")
    for row in overview.iter_rows():
        for c in row:
            if isinstance(c.value, str):
                m = pat.match(c.value)
                if m:
                    c.value = "=" + total(m.group(1), m.group(2))
    return created
